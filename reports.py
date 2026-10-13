"""Отчёты по нашей копии МойСклад (только чтение): прибыль по товарам и клиентам, обороты товаров, кассовые смены.

- продажи — проведённые отгрузки и чеки минус возвраты покупателей и по чекам (позиции документов, сумма со скидкой);
- себестоимость — количество × закупочная цена из карточки товара (приближение: МойСклад считает по партиям,
  поэтому прибыль здесь может немного отличаться от его отчёта «Прибыльность»);
- период — по дате документа в Алматы, не больше окна зеркала (DOC_DAYS);
- обороты — приход (приёмки, оприходования, возвраты покупателей) и расход (отгрузки, чеки, списания, возвраты поставщикам)
  по товарам за период и текущий остаток;
- кассовые смены — смены МойСклад (выручка наличными и безналом — как в МойСклад) и их чеки и возвраты из нашей копии
  (чек относится к смене по ссылке МойСклад; ещё не перечитанный — по складу и времени между открытием и закрытием). Только просмотр: без ОФД и без записи в МойСклад."""
import time
from datetime import datetime, timedelta

import mirror
import order_hook as oh

SALE = ("demand", "retaildemand")
RET = ("salesreturn", "retailsalesreturn")
IN_DOCS = ("supply", "enter", "salesreturn", "retailsalesreturn")
OUT_DOCS = ("demand", "retaildemand", "loss", "purchasereturn")


def _bounds(days):
    """Последние days дней по Алматы (сегодня включительно) → границы moment в московском времени МойСклад."""
    days = max(1, min(int(days or 30), mirror.DOC_DAYS))
    today = datetime.now(oh.ALMATY).date()
    start = datetime(*(today - timedelta(days=days - 1)).timetuple()[:3], tzinfo=oh.ALMATY)
    end = datetime(*(today + timedelta(days=1)).timetuple()[:3], tzinfo=oh.ALMATY)
    f = lambda t: t.astimezone(mirror.MS_TZ).strftime("%Y-%m-%d %H:%M:%S")
    return days, f(start), f(end)


def _ph(seq):
    return ", ".join(["%s"] * len(seq))


def profit(days=30, by="product", limit=50):
    """Прибыль за период по товарам (by=product) или клиентам (by=agent): выручка, себестоимость, прибыль, маржа."""
    days, a, b = _bounds(days)
    key = "l.product_id" if by == "product" else "x.agent_id"
    agg = {}
    with mirror.db() as d:
        for k, t, q, sm, buy in d.run(f"SELECT {key}, x.type, l.qty, l.sum, COALESCE(p.buy_price, 0)"
                                      f" FROM ms_doc x JOIN ms_doc_line l ON l.doc_id=x.id LEFT JOIN ms_product p ON p.id=l.product_id"
                                      f" WHERE x.deleted=0 AND x.applicable=1 AND x.moment>=%s AND x.moment<%s AND x.type IN ({_ph(SALE + RET)})",
                                      (a, b, *SALE, *RET), many=True) or []:
            sg = -1 if t in RET else 1
            r = agg.setdefault(k, [0, 0, 0])
            r[0] += sg * (q or 0)
            r[1] += sg * (sm or 0)
            r[2] += sg * (q or 0) * (buy or 0)
        rows = [(k, *v) for k, v in agg.items()]
        ids = [r[0] for r in rows if r[0]]
        if by == "product":
            names = {i: (n, c) for i, n, c in d.run(f"SELECT id, name, code FROM ms_product WHERE id IN ({_ph(ids)})", ids, many=True) or []} if ids else {}
        else:
            names = {i: (n, "") for i, n in d.run(f"SELECT id, name FROM ms_agent WHERE id IN ({_ph(ids)})", ids, many=True) or []} if ids else {}
    out = []
    for k, q, rev, cost in rows:
        rev, cost = round(rev or 0), round(cost or 0)
        n, c = names.get(k, ("(удалён)" if k else "Без контрагента", ""))
        out.append({"id": k or "", "name": n or "—", "code": c or "", "qty": round(q or 0, 3), "revenue": rev, "cost": cost,
                    "profit": rev - cost, "margin": round((rev - cost) * 100 / rev, 1) if rev else 0})
    out.sort(key=lambda r: -r["revenue"])
    tot = {k: sum(r[k] for r in out) for k in ("revenue", "cost", "profit")}
    tot["margin"] = round(tot["profit"] * 100 / tot["revenue"], 1) if tot["revenue"] else 0
    tot["noCost"] = sum(1 for r in out if r["revenue"] and not r["cost"]) if by == "product" else 0
    return {"days": days, "by": by, "total": tot, "count": len(out), "rows": out[:limit]}


def turnover(days=30, q="", limit=50):
    """Обороты товаров за период: приход, расход (шт.) и текущий остаток — по всем складам."""
    days, a, b = _bounds(days)
    q = (q or "").strip()
    with mirror.db() as d:
        rows = d.run(f"SELECT l.product_id, SUM(CASE WHEN x.type IN ({_ph(IN_DOCS)}) THEN l.qty ELSE 0 END),"
                     f" SUM(CASE WHEN x.type IN ({_ph(OUT_DOCS)}) THEN l.qty ELSE 0 END)"
                     f" FROM ms_doc x JOIN ms_doc_line l ON l.doc_id=x.id"
                     f" WHERE x.deleted=0 AND x.applicable=1 AND x.moment>=%s AND x.moment<%s AND x.type IN ({_ph(IN_DOCS + OUT_DOCS)})"
                     f" GROUP BY l.product_id", (*IN_DOCS, *OUT_DOCS, a, b, *IN_DOCS, *OUT_DOCS), many=True) or []
        ids = [r[0] for r in rows if r[0]]
        info = {i: (n, c) for i, n, c in d.run(f"SELECT id, name, code FROM ms_product WHERE id IN ({_ph(ids)})", ids, many=True) or []} if ids else {}
        stock = dict(d.run(f"SELECT product_id, SUM(stock) FROM ms_stock WHERE product_id IN ({_ph(ids)}) GROUP BY product_id", ids, many=True) or []) if ids else {}
    out = []
    for pid, i, o in rows:
        n, c = info.get(pid, ("(удалён)", ""))
        if q and q.lower() not in (n or "").lower() and q not in (c or ""):
            continue
        out.append({"id": pid, "name": n or "—", "code": c or "", "in": round(i or 0, 3), "out": round(o or 0, 3), "stock": round(stock.get(pid, 0) or 0, 3)})
    out.sort(key=lambda r: -(r["out"] + r["in"]))
    return {"days": days, "count": len(out), "totalIn": round(sum(r["in"] for r in out), 3), "totalOut": round(sum(r["out"] for r in out), 3),
            "rows": out[:limit]}


def xlsx(kind="product", days=30, q=""):
    """Отчёт целиком (без ограничения строк) в Excel: прибыль по товарам / клиентам или обороты."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    if kind == "turn":
        r = turnover(days, q, limit=100000)
        ws.title = "Обороты"
        head = ["Товар", "Код", "Приход, шт.", "Расход, шт.", "Остаток, шт."]
        rows = [[x["name"], x["code"], x["in"], x["out"], x["stock"]] for x in r["rows"]]
        total = ["Итого", "", r["totalIn"], r["totalOut"], ""]
    else:
        by = "agent" if kind == "agent" else "product"
        r = profit(days, by, limit=100000)
        ws.title = "Прибыль по клиентам" if by == "agent" else "Прибыль по товарам"
        head = (["Клиент"] if by == "agent" else ["Товар", "Код", "Кол-во, шт."]) + ["Выручка, ₸", "Себестоимость, ₸", "Прибыль, ₸", "Маржа, %"]
        rows = [([x["name"]] if by == "agent" else [x["name"], x["code"], x["qty"]]) + [x["revenue"], x["cost"], x["profit"], x["margin"]]
                for x in r["rows"]]
        t = r["total"]
        total = (["Итого"] if by == "agent" else ["Итого", "", ""]) + [t["revenue"], t["cost"], t["profit"], t["margin"]]
    ws.append([f"AMURA — {ws.title} за {r['days']} дн. (по данным МойСклад в системе AMURA)"])
    ws["A1"].font = Font(bold=True)
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.append(total)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 48
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RETAIL = ("retaildemand", "retailsalesreturn")


_pts = {}


def _points():
    """Названия точек продаж (кэш на час, только чтение). МойСклад не ответил — без названий."""
    c = _pts.get("v")
    if not c or time.time() - c[0] > 3600:
        try:
            rows = oh.ms("GET", "/entity/retailstore", params={"limit": 100}, timeout=20).get("rows", [])
            c = (time.time(), {r.get("id"): r.get("name") or "" for r in rows})
        except Exception:
            c = (time.time() - 3000, (c or (0, {}))[1])          # повторим через 10 минут
        _pts["v"] = c
    return c[1]


def _shift_row(r):
    i, num, mo, cl, store, point, cash, nocash, gc, gn, com = r
    return {"id": i, "number": num or "", "open": mo or "", "close": cl or "", "storeId": store or "", "pointId": point or "", "cash": round(cash or 0),
            "noCash": round(nocash or 0), "total": round((cash or 0) + (nocash or 0)), "gotCash": round(gc or 0),
            "gotNoCash": round(gn or 0), "commission": round(com or 0)}


def _in_shift(s, doc_moment, doc_store, doc_shift):
    """Чек относится к смене по ссылке МойСклад (retailShift). Чек, ещё не перечитанный со сменой, — по складу и времени
    между открытием и закрытием (открытая смена — до сих пор); у двух точек на одном складе так можно ошибиться."""
    if doc_shift:
        return doc_shift == s["id"]
    return (not s["storeId"] or doc_store == s["storeId"]) and s["open"] <= doc_moment and (not s["close"] or doc_moment <= s["close"])


def shifts(days=30, limit=100):
    """Кассовые смены за период (по открытию, Алматы): выручка наличными / безналом по МойСклад и чеки из нашей копии."""
    days, a, b = _bounds(days)
    cols = "id, number, moment, close_at, store_id, point, cash, nocash, got_cash, got_nocash, commission"
    with mirror.db() as d:
        rows = [_shift_row(r) for r in d.run(f"SELECT {cols} FROM ms_shift WHERE deleted=0 AND moment>=%s AND moment<%s ORDER BY moment DESC",
                                             (a, b), many=True) or []]
        docs = d.run(f"SELECT type, moment, store_id, sum, shift_id FROM ms_doc WHERE deleted=0 AND applicable=1 AND type IN ({_ph(RETAIL)}) AND moment>=%s",
                     (*RETAIL, a), many=True) or []
        stores = dict(d.run("SELECT id, name FROM ms_store", many=True) or [])
    pts = _points() if len({s["pointId"] for s in rows}) > 1 else {}       # название точки — только если их несколько
    for s in rows:
        s["store"] = stores.get(s["storeId"], "")
        s["point"] = pts.get(s["pointId"], "")
        s["receipts"] = s["returns"] = s["returnsSum"] = 0
        for t, mo, st, sm, sh in docs:
            if _in_shift(s, mo or "", st or "", sh or ""):
                if t == "retaildemand":
                    s["receipts"] += 1
                else:
                    s["returns"] += 1
                    s["returnsSum"] += round(sm or 0)
    tot = {"cash": sum(s["cash"] for s in rows), "noCash": sum(s["noCash"] for s in rows), "receipts": sum(s["receipts"] for s in rows)}
    tot["total"] = tot["cash"] + tot["noCash"]
    tot["avg"] = round(tot["total"] / len(rows)) if rows else 0
    return {"days": days, "count": len(rows), "open": sum(1 for s in rows if not s["close"]), "total": tot, "rows": rows[:limit]}


def shift(shift_id):
    """Одна смена и её чеки и возвраты (из нашей копии МойСклад). Нет смены — None."""
    cols = "id, number, moment, close_at, store_id, point, cash, nocash, got_cash, got_nocash, commission"
    with mirror.db() as d:
        r = d.run(f"SELECT {cols} FROM ms_shift WHERE id=%s AND deleted=0", (shift_id,), one=True)
        if not r:
            return None
        s = _shift_row(r)
        win = "x.moment>=%s" + (" AND x.moment<=%s" if s["close"] else "") + (" AND x.store_id=%s" if s["storeId"] else "")
        args = [*RETAIL, s["id"], s["open"]] + ([s["close"]] if s["close"] else []) + ([s["storeId"]] if s["storeId"] else [])
        docs = d.run(f"SELECT x.id, x.type, x.number, x.moment, x.sum, a.name FROM ms_doc x LEFT JOIN ms_agent a ON a.id=x.agent_id"
                     f" WHERE x.deleted=0 AND x.applicable=1 AND x.type IN ({_ph(RETAIL)})"
                     f" AND (x.shift_id=%s OR (COALESCE(x.shift_id, '')='' AND {win})) ORDER BY x.moment DESC", args, many=True) or []
        st = d.run("SELECT name FROM ms_store WHERE id=%s", (s["storeId"],), one=True) if s["storeId"] else None
        lines = {}
        ids = [x[0] for x in docs]
        if ids:
            for doc_id, n in d.run(f"SELECT doc_id, COUNT(*) FROM ms_doc_line WHERE doc_id IN ({_ph(ids)}) GROUP BY doc_id", ids, many=True) or []:
                lines[doc_id] = n
    s["store"] = (st and st[0]) or ""
    s["point"] = _points().get(s["pointId"], "")
    s["docs"] = [{"id": i, "type": t, "title": "Возврат" if t == "retailsalesreturn" else "Чек", "number": n or "", "at": mo or "",
                  "sum": round(sm or 0), "agent": ag or "", "lines": lines.get(i, 0)} for i, t, n, mo, sm, ag in docs]
    s["receipts"] = sum(1 for x in s["docs"] if x["type"] == "retaildemand")
    s["receiptsSum"] = sum(x["sum"] for x in s["docs"] if x["type"] == "retaildemand")
    s["returnsSum"] = sum(x["sum"] for x in s["docs"] if x["type"] != "retaildemand")
    return s
