"""Модуль «Источники»: откуда пришли покупатели (авторы UGC, блогеры, реклама) и сколько купили.
Ссылка с меткой: https://amura.kz/?ref=CODE (сайт запоминает метку на 30 дней). Промокод источника — поле в корзине
только на розничном сайте (на оптовом его нет — решение владельца 13.10);
скидка по промокоду только если владелец задал % > 0, цену пересчитывает сервер. Без метки и промокода заказ идёт как раньше.
Таблицы: src_source (источники), src_event (переходы и «в корзину» по дням), src_order (заказы с сайта и их источник)."""
import json
import re
import time
from datetime import datetime, timedelta

from flask import Blueprint, jsonify, request

import inbox
import modules
import order_hook as oh

bp = Blueprint("sources", __name__)
bp.after_request(oh.cors)
CODE_RE = re.compile(r"^[a-z0-9][a-z0-9_-]{1,31}$")
_ready = [False]


def db():
    d = inbox.db()
    if not _ready[0]:
        d.run("CREATE TABLE IF NOT EXISTS src_source (code TEXT PRIMARY KEY, name TEXT, cost DOUBLE PRECISION DEFAULT 0, promo TEXT DEFAULT '', "
              "discount DOUBLE PRECISION DEFAULT 0, archived INTEGER DEFAULT 0, at DOUBLE PRECISION)")
        d.run("CREATE TABLE IF NOT EXISTS src_event (day TEXT, code TEXT, kind TEXT, retail INTEGER, n INTEGER, PRIMARY KEY (day, code, kind, retail))")
        d.run("CREATE TABLE IF NOT EXISTS src_order (order_id TEXT PRIMARY KEY, number TEXT, code TEXT, goods DOUBLE PRECISION, retail INTEGER, "
              "lines TEXT, at DOUBLE PRECISION)")
        d.c.commit()
        _ready[0] = True
    return d


def on():
    return modules.enabled("sources")


def norm(code):
    code = str(code or "").strip().lower()
    return code if CODE_RE.match(code) else ""


def _row(r):
    return {"code": r[0], "name": r[1], "cost": round(r[2] or 0), "promo": r[3] or "", "discount": r[4] or 0, "archived": bool(r[5])}


def all_sources():
    with db() as d:
        return [_row(r) for r in d.run("SELECT code, name, cost, promo, discount, archived, at FROM src_source ORDER BY at", many=True) or []]


def _lookup():
    """Активные источники (кэш 30 с): {code: src} и {PROMO: src}."""
    def get():
        s = [x for x in all_sources() if not x["archived"]]
        return {x["code"]: x for x in s}, {x["promo"].upper(): x for x in s if x["promo"]}
    return oh.cached("src_lookup", 30, get)


def reset():
    oh._cache.pop("src_lookup", None)


def by_code(code):
    return _lookup()[0].get(norm(code))


def by_promo(promo):
    promo = str(promo or "").strip().upper()[:32]
    return _lookup()[1].get(promo) if promo else None


def promo_available():
    try:
        return on() and bool(_lookup()[1])
    except Exception:
        return False


def save(b):
    """Добавить или изменить источник (код не меняется)."""
    code = norm(b.get("code"))
    name = str(b.get("name", "")).strip()[:80]
    if not code:
        raise ValueError("Код — латиница, цифры, «-» или «_», от 2 до 32 знаков")
    if not name:
        raise ValueError("Укажите имя источника")
    try:
        cost = max(0.0, float(b.get("cost") or 0))
        disc = min(90.0, max(0.0, float(b.get("discount") or 0)))
    except (TypeError, ValueError):
        raise ValueError("Стоимость и скидка — числа")
    promo = re.sub(r"[^A-Za-z0-9А-Яа-яЁё_-]", "", str(b.get("promo", "")))[:32].upper()
    with db() as d:
        if promo and d.run("SELECT code FROM src_source WHERE promo=%s AND code<>%s", (promo, code), one=True):
            raise ValueError("Такой промокод уже есть у другого источника")
        if d.run("SELECT code FROM src_source WHERE code=%s", (code,), one=True):
            d.run("UPDATE src_source SET name=%s, cost=%s, promo=%s, discount=%s, archived=%s WHERE code=%s",
                  (name, cost, promo, disc, 1 if b.get("archived") else 0, code))
        else:
            d.run("INSERT INTO src_source (code, name, cost, promo, discount, archived, at) VALUES (%s, %s, %s, %s, %s, 0, %s)",
                  (code, name, cost, promo, disc, time.time()))
    reset()


def _today():
    return datetime.now(oh.ALMATY).strftime("%Y-%m-%d")


def count(code, kind, retail):
    with db() as d:
        k = (_today(), code, kind, 1 if retail else 0)
        if d.run("SELECT n FROM src_event WHERE day=%s AND code=%s AND kind=%s AND retail=%s", k, one=True):
            d.run("UPDATE src_event SET n=n+1 WHERE day=%s AND code=%s AND kind=%s AND retail=%s", k)
        else:
            d.run("INSERT INTO src_event (day, code, kind, retail, n) VALUES (%s, %s, %s, %s, 1)", k)


# ---------- заказ: источник, промокод, запись ----------
def for_order(d):
    """Источник заказа с сайта: верный промокод важнее метки. Возвращает (src | None, скидка %)."""
    if not on():
        return None, 0
    s = by_promo(d.get("promo")) if str(d.get("mode", "")) == "retail" else None   # промокод — только на розничном сайте
    if s:
        return s, (s["discount"] if s["discount"] > 0 else 0)
    return by_code(d.get("ref")), 0


def apply_discount(price, disc):
    return max(1, int(price * (100 - disc) / 100 + 0.5)) if disc > 0 else price   # как Math.round на сайте


def record_order(order_id, number, code, lines, retail):
    """Каждый заказ с сайта (и без источника — для строки «без источника»). Сбой не ломает заказ."""
    if not on():
        return
    try:
        goods = sum(l["qty"] * l["price"] for l in lines)
        ls = json.dumps([[l["id"], l["name"][:80], l["qty"], l["qty"] * l["price"]] for l in lines], ensure_ascii=False)
        with db() as d:
            d.run("INSERT INTO src_order (order_id, number, code, goods, retail, lines, at) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                  (order_id, str(number), code or "", goods, 1 if retail else 0, ls, time.time()))
    except Exception as e:
        print("Источники: заказ не записан:", str(e)[:200], flush=True)


# ---------- сайт: переходы и проверка промокода ----------
def _ip():
    return request.headers.get("X-Forwarded-For", request.remote_addr or "").split(",")[0].strip()


@bp.route("/ref/visit", methods=["POST", "OPTIONS"])
def visit():
    """Переход по ссылке (раз в день на устройство) или первое «в корзину» за визит. Без личных данных."""
    if request.method == "OPTIONS":
        return "", 204
    if not on():
        return jsonify(ok=False), 404
    if oh.too_many("ref:" + _ip(), limit=60, window=3600):
        return jsonify(ok=False), 429
    b = request.get_json(silent=True) or {}
    s = by_code(b.get("ref"))
    if not s:
        return jsonify(ok=False), 404
    kind = "cart" if b.get("kind") == "cart" else "visit"
    try:
        count(s["code"], kind, bool(b.get("retail")))
    except Exception as e:
        print("Источники: переход не записан:", str(e)[:200], flush=True)
    return jsonify(ok=True)


@bp.route("/ref/promo", methods=["GET", "OPTIONS"])
def promo_check():
    if request.method == "OPTIONS":
        return "", 204
    if not on() or request.args.get("retail") != "1" or oh.too_many("promo:" + _ip(), limit=30, window=600):   # только розница
        return jsonify(ok=False, error="Промокод не найден"), 404
    s = by_promo(request.args.get("code"))
    if not s:
        return jsonify(ok=False, error="Промокод не найден"), 404
    return jsonify(ok=True, code=s["promo"], discount=s["discount"] if s["discount"] > 0 else 0)


# ---------- отчёт ----------
def _cancelled_ids(order_ids):
    """Заказы, отменённые или удалённые в МойСклад (по нашей копии): статус со словом «отмен» или «возврат»."""
    if not order_ids:
        return set()
    try:
        md = oh.cached("src_states", 600, lambda: oh.ms("GET", "/entity/customerorder/metadata", timeout=12))
        bad = {s["id"] for s in md.get("states", []) if re.search(r"отмен|возврат", s.get("name", ""), re.I)}
        import mirror
        out = set()
        with mirror.db() as d:
            ids = list(order_ids)
            for k in range(0, len(ids), 500):
                part = ids[k:k + 500]
                q = ",".join(["%s"] * len(part))
                for oid, st, dl in d.run(f"SELECT id, state_id, deleted FROM ms_doc WHERE type='customerorder' AND id IN ({q})", part, many=True) or []:
                    if dl or st in bad:
                        out.add(oid)
        return out
    except Exception as e:
        print("Источники: статусы заказов не получены:", str(e)[:200], flush=True)
        return set()


def report(days=30):
    days = days if days in (7, 30, 90) else 30
    since = time.time() - days * 86400
    day0 = (datetime.now(oh.ALMATY) - timedelta(days=days - 1)).strftime("%Y-%m-%d")
    srcs = {s["code"]: s for s in all_sources()}
    with db() as d:
        ev = d.run("SELECT code, kind, SUM(n) FROM src_event WHERE day>=%s GROUP BY code, kind", (day0,), many=True) or []
        orders = d.run("SELECT order_id, code, goods, lines FROM src_order WHERE at>=%s", (since,), many=True) or []
    bad = _cancelled_ids([o[0] for o in orders])
    rows = {}

    def row(code):
        if code not in rows:
            s = srcs.get(code) or {}
            rows[code] = {"code": code, "name": s.get("name") or ("Без источника" if not code else code), "cost": s.get("cost", 0),
                          "visits": 0, "carts": 0, "orders": 0, "revenue": 0, "cancelled": 0, "_p": {}}
        return rows[code]
    for s in srcs.values():
        if not s["archived"]:
            row(s["code"])
    for code, kind, n in ev:
        row(code)["visits" if kind == "visit" else "carts"] += int(n or 0)
    for oid, code, goods, ls in orders:
        r = row(code or "")
        if oid in bad:
            r["cancelled"] += 1
            continue
        r["orders"] += 1
        r["revenue"] += goods or 0
        for pid, name, qty, sm in json.loads(ls or "[]"):
            p = r["_p"].setdefault(pid, {"name": name, "qty": 0, "sum": 0})
            p["qty"] += qty
            p["sum"] += sm
    out = []
    for r in rows.values():
        r["revenue"] = round(r["revenue"])
        r["conv"] = round(r["orders"] * 100 / r["visits"], 1) if r["visits"] else 0
        r["avg"] = round(r["revenue"] / r["orders"]) if r["orders"] else 0
        r["payback"] = round(r["revenue"] / r["cost"], 2) if r["cost"] else None
        r["top"] = [{"name": p["name"], "qty": p["qty"], "sum": round(p["sum"])} for p in sorted(r.pop("_p").values(), key=lambda p: -p["sum"])[:3]]
        out.append(r)
    out.sort(key=lambda r: (r["code"] == "", -r["revenue"], -r["visits"]))
    tot = {k: sum(r[k] for r in out if r["code"]) for k in ("visits", "carts", "orders", "revenue", "cost")}
    return {"days": days, "rows": out, "total": tot}


def xlsx(days=30):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = report(days)
    wb = Workbook()
    ws = wb.active
    ws.title = "Источники"
    ws.append([f"AMURA — источники покупателей за {r['days']} дн. (выручка — заказы с сайта без отменённых)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Источник", "Код", "Переходы", "В корзину", "Заказы", "Выручка, ₸", "Конверсия, %", "Средний чек, ₸",
               "Стоимость, ₸", "Окупаемость", "Отменено", "Лучшие товары"])
    for c in ws[2]:
        c.font = Font(bold=True)
    for x in r["rows"]:
        ws.append([x["name"], x["code"], x["visits"], x["carts"], x["orders"], x["revenue"], x["conv"], x["avg"], x["cost"] or "",
                   x["payback"] if x["payback"] is not None else "", x["cancelled"], "; ".join(f"{p['name']} — {p['qty']} шт." for p in x["top"])])
    for col, w in zip("ABCDEFGHIJKL", (28, 14, 10, 10, 9, 13, 12, 13, 12, 12, 10, 60)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
