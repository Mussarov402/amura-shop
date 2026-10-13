"""Отчёты по зеркалу: прибыль по товарам и клиентам (продажи − возвраты, себестоимость по закупочной цене), обороты."""
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import test_mirror as tm  # noqa: E402,F401
import admin  # noqa: E402
import app  # noqa: E402
import mirror  # noqa: E402
import order_hook as oh  # noqa: E402


def now_ms():
    return datetime.now(mirror.MS_TZ).strftime("%Y-%m-%d %H:%M:%S")


class ReportsTest(unittest.TestCase):
    def setUp(self):
        oh.ORDER_SECRET = b"x"
        self.c = app.app.test_client()
        self.h = {"Authorization": "Bearer " + admin.make_admin_token()}
        t = now_ms()
        with mirror.db() as d:
            for x in ("ms_doc", "ms_doc_line", "ms_product", "ms_agent", "ms_stock"):
                d.run(f"DELETE FROM {x}")
            d.run("DELETE FROM setting WHERE key='mod_finance'")
            d.run("INSERT INTO ms_product (id, kind, code, name, buy_price, archived, deleted) VALUES ('p1', 'product', '1', 'Krem', 600, 0, 0)")
            d.run("INSERT INTO ms_product (id, kind, code, name, buy_price, archived, deleted) VALUES ('p2', 'product', '2', 'Shampoo', 0, 0, 0)")
            d.run("INSERT INTO ms_agent (id, name, archived, deleted) VALUES ('k1', 'Ertos', 0, 0)")
            docs = [("d1", "demand", "k1", 1, t), ("r1", "retaildemand", "", 1, t), ("sr1", "salesreturn", "k1", 1, t),
                    ("d0", "demand", "k1", 0, t), ("old", "demand", "k1", 1, "2020-01-01 10:00:00"), ("sp1", "supply", "", 1, t), ("ls1", "loss", "", 1, t)]
            for i, ty, ag, ap, mo in docs:
                d.run("INSERT INTO ms_doc (id, type, number, moment, agent_id, sum, applicable, deleted) VALUES (%s, %s, '1', %s, %s, 0, %s, 0)", (i, ty, mo, ag, ap))
            for doc, pid, q, s in [("d1", "p1", 3, 3000), ("d1", "p2", 1, 500), ("r1", "p1", 1, 1100), ("sr1", "p1", 1, 1000),
                                   ("d0", "p1", 100, 100000), ("old", "p1", 100, 100000), ("sp1", "p1", 10, 6000), ("ls1", "p2", 2, 0)]:
                d.run("INSERT INTO ms_doc_line (doc_id, pos_id, product_id, qty, price, discount, sum) VALUES (%s, %s, %s, %s, 0, 0, %s)", (doc, doc + pid, pid, q, s))
            d.run("INSERT INTO ms_stock (product_id, store_id, stock, reserve, synced) VALUES ('p1', 'sA', 7, 0, 0)")

    def get(self, path):
        return self.c.get("/admin/api/reports/" + path, headers=self.h).get_json()

    def test_profit_by_product(self):
        r = self.get("profit?days=30")
        rows = {x["id"]: x for x in r["rows"]}
        # p1: 3 + 1 (чек) − 1 (возврат) = 3 шт., выручка 3000 + 1100 − 1000 = 3100, себестоимость 3 × 600 = 1800
        self.assertEqual((rows["p1"]["qty"], rows["p1"]["revenue"], rows["p1"]["cost"], rows["p1"]["profit"]), (3, 3100, 1800, 1300))
        self.assertEqual((rows["p2"]["revenue"], rows["p2"]["cost"]), (500, 0))
        self.assertEqual(r["total"]["revenue"], 3600)                     # непроведённая и старая отгрузки не считаются
        self.assertEqual(r["total"]["noCost"], 1)                         # у одного товара нет закупочной цены
        self.assertEqual(r["rows"][0]["id"], "p1")

    def test_profit_by_client(self):
        rows = {x["name"]: x for x in self.get("profit?days=30&by=agent")["rows"]}
        self.assertEqual(rows["Ertos"]["revenue"], 2500)                   # 3000 + 500 − 1000
        self.assertEqual(rows["Без контрагента"]["revenue"], 1100)

    def test_turnover(self):
        rows = {x["id"]: x for x in self.get("turnover?days=30")["rows"]}
        self.assertEqual((rows["p1"]["in"], rows["p1"]["out"], rows["p1"]["stock"]), (11, 4, 7))   # приёмка 10 + возврат 1; отгрузка 3 + чек 1
        self.assertEqual((rows["p2"]["in"], rows["p2"]["out"]), (0, 3))                            # отгрузка 1 + списание 2
        self.assertEqual([x["id"] for x in self.get("turnover?days=30&q=Sham")["rows"]], ["p2"])

    def test_excel_export(self):
        import io
        from openpyxl import load_workbook
        for kind, first in (("product", "Krem"), ("agent", "Ertos"), ("turn", "Shampoo")):
            r = self.c.post("/admin/api/reports/export", headers=self.h, json={"kind": kind, "days": 30}).get_json()
            f = self.c.get(r["url"])
            self.assertEqual(f.status_code, 200)
            ws = load_workbook(io.BytesIO(f.data)).active
            names = [c.value for c in ws["A"]][2:-1]
            self.assertIn(first, names)
            self.assertEqual(ws.cell(ws.max_row, 1).value, "Итого")
        ws = load_workbook(io.BytesIO(self.c.get(self.c.post("/admin/api/reports/export", headers=self.h, json={"kind": "product"}).get_json()["url"]).data)).active
        self.assertEqual([c.value for c in ws[ws.max_row]][3:6], [3600, 1800, 1800])           # итог: выручка, себестоимость, прибыль
        self.assertEqual(self.c.get(r["url"].split("?")[0] + "?t=bad").status_code, 404)
        self.assertEqual(self.c.post("/admin/api/reports/export", json={"kind": "turn"}).status_code, 401)

    def test_owner_and_module(self):
        self.assertEqual(self.c.get("/admin/api/reports/profit").status_code, 401)
        self.c.put("/admin/api/modules", headers=self.h, json={"id": "finance", "on": False})
        self.assertEqual(self.c.get("/admin/api/reports/turnover", headers=self.h).status_code, 404)


if __name__ == "__main__":
    unittest.main()
