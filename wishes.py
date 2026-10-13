"""Модуль «Пожелания клиентов»: какие бренды и товары хотят видеть покупатели и что улучшить.
Анкета на сайте (amura.kz/#opros, opt.amura.kz/#opros) → POST /wish; менеджер переносит сообщение из чата кнопкой «В пожелания».
Скидки и промокодов за ответ нет. Таблица wish в нашей базе."""
import json
import re
import time

from flask import Blueprint, jsonify, request

import inbox
import modules
import order_hook as oh

bp = Blueprint("wishes", __name__)
bp.after_request(oh.cors)
STATUSES = ("new", "done", "answered")          # новое / учли / ответили
_ready = [False]


def db():
    d = inbox.db()
    if not _ready[0]:
        pk = "SERIAL PRIMARY KEY" if inbox.PG else "INTEGER PRIMARY KEY AUTOINCREMENT"
        d.run(f"CREATE TABLE IF NOT EXISTS wish (id {pk}, at DOUBLE PRECISION, site TEXT, ref TEXT, brands TEXT, want TEXT, improve TEXT, "
              "rating INTEGER, name TEXT, phone TEXT, status TEXT DEFAULT 'new', conv_id INTEGER)")
        d.c.commit()
        _ready[0] = True
    return d


def on():
    return modules.enabled("wishes")


def _catalog_brands():
    """Бренды каталога: {нижний регистр: (название, остаток, товаров)} — из живого каталога (и скрытых), без запроса в МойСклад сверх кэша."""
    out = {}
    try:
        v = oh._cache.get("live")
        data = v[1] if v else oh.live()
        for i in (data.get("items") or []) + (data.get("hiddenItems") or []):
            b = (i.get("brand") or "").strip()
            if b:
                n, q, c = out.get(b.lower(), (b, 0, 0))
                out[b.lower()] = (n, q + int(i.get("qty") or 0), c + 1)
    except Exception as e:
        print("Пожелания: каталог не получен:", str(e)[:200], flush=True)
    return out


def _brands_in(text, known):
    """Бренды каталога, упомянутые в тексте (для сообщений из чата)."""
    low = " " + re.sub(r"\s+", " ", text.lower()) + " "
    return [n for k, (n, _, _) in known.items() if len(k) >= 3 and re.search(r"(?<![\w])" + re.escape(k) + r"(?![\w])", low)]


def _clean_brands(v):
    out = []
    for b in (v if isinstance(v, list) else [])[:10]:
        b = re.sub(r"\s+", " ", str(b)).strip()[:40]
        if b and b.lower() not in [x.lower() for x in out]:
            out.append(b)
    return out


def add(site, want="", improve="", brands=None, rating=None, name="", phone="", ref="", conv_id=None):
    with db() as d:
        return d.run("INSERT INTO wish (at, site, ref, brands, want, improve, rating, name, phone, status, conv_id) "
                     "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'new', %s)",
                     (time.time(), site, ref, json.dumps(_clean_brands(brands), ensure_ascii=False), want, improve, rating, name, phone, conv_id), ins=True)


@bp.route("/wish", methods=["GET", "POST", "OPTIONS"])
def wish_post():
    """Ответ на анкету (GET — открыта ли анкета). Лимит по IP, имя и телефон — только если покупатель сам указал."""
    if request.method == "OPTIONS":
        return "", 204
    if request.method == "GET":
        return jsonify(ok=on())
    if not on():
        return jsonify(ok=False, error="Анкета сейчас недоступна"), 404
    ip = request.headers.get("X-Forwarded-For", request.remote_addr or "").split(",")[0].strip()
    if oh.too_many("wish:" + ip, limit=5, window=3600):
        return jsonify(ok=False, error="Спасибо, ответ уже получен"), 429
    b = request.get_json(silent=True) or {}
    want = str(b.get("want", "")).strip()[:1000]
    improve = str(b.get("improve", "")).strip()[:1000]
    brands = _clean_brands(b.get("brands"))
    try:
        rating = int(b.get("rating") or 0)
    except (TypeError, ValueError):
        rating = 0
    rating = rating if 1 <= rating <= 5 else None
    if not (want or improve or brands or rating):
        return jsonify(ok=False, error="Напишите хотя бы одно пожелание"), 400
    phone = re.sub(r"\D", "", str(b.get("phone", "")))[:15]
    ref = str(b.get("ref", "")).strip().lower()[:32]
    ref = ref if re.fullmatch(r"[a-z0-9_-]*", ref) else ""
    add("rtl" if b.get("retail") else "opt", want, improve, brands, rating, str(b.get("name", "")).strip()[:60], phone, ref)
    return jsonify(ok=True)


def from_message(msg_id):
    """Входящее сообщение из «Сообщений» → пожелание (бренды каталога из текста распознаются сами)."""
    with inbox.db() as d:
        m = d.run("SELECT conv_id, text, role FROM msg WHERE id=%s", (int(msg_id),), one=True)
        c = d.run("SELECT name, chat_id FROM conv WHERE id=%s", (m[0],), one=True) if m else None
    if not m or m[2] != "client" or not (m[1] or "").strip():
        raise ValueError("Это не сообщение клиента с текстом")
    with db() as d:
        if d.run("SELECT id FROM wish WHERE conv_id=%s AND want=%s", (m[0], m[1][:1000]), one=True):
            raise ValueError("Уже в пожеланиях")
    chat = (c[1] if c else "") or ""
    site = "ig" if chat.startswith("ig:") else "wa" if chat.startswith("wa:") else "tg" if chat[:1].isdigit() or chat.startswith("-") else "chat"
    return add(site, m[1][:1000], "", _brands_in(m[1], _catalog_brands()), None, (c[0] if c else "")[:60], "", "", m[0])


def set_status(wid, status):
    if status not in STATUSES:
        raise ValueError("Неизвестная отметка")
    with db() as d:
        d.run("UPDATE wish SET status=%s WHERE id=%s", (status, int(wid)))


SITE_NAMES = {"opt": "Анкета · опт", "rtl": "Анкета · розница", "ig": "Instagram", "wa": "WhatsApp", "tg": "Telegram", "chat": "Чат"}


def report(days=30):
    days = days if days in (7, 30, 90) else 30
    with db() as d:
        rows = d.run("SELECT id, at, site, ref, brands, want, improve, rating, name, phone, status FROM wish WHERE at>=%s ORDER BY at DESC",
                     (time.time() - days * 86400,), many=True) or []
    known = _catalog_brands()
    top, items, rates = {}, [], []
    for wid, at, site, ref, brands, want, improve, rating, name, phone, status in rows:
        bl = json.loads(brands or "[]")
        for b in bl:
            k = b.lower()
            t = top.setdefault(k, {"brand": known.get(k, (b,))[0], "count": 0})
            t["count"] += 1
        if rating:
            rates.append(rating)
        items.append({"id": wid, "at": at, "site": site, "siteName": SITE_NAMES.get(site, site), "ref": ref or "", "brands": bl,
                      "want": want or "", "improve": improve or "", "rating": rating, "name": name or "", "phone": phone or "", "status": status or "new"})
    for k, t in top.items():
        n, qty, cnt = known.get(k, (t["brand"], 0, 0))
        t.update(inCatalog=cnt > 0, qty=qty, products=cnt)
    return {"days": days, "count": len(items), "rating": round(sum(rates) / len(rates), 1) if rates else None, "ratings": len(rates),
            "brands": sorted(top.values(), key=lambda t: (-t["count"], t["brand"].lower()))[:50], "items": items[:300],
            "open": sum(1 for i in items if i["status"] == "new")}


def brand_hints():
    """Подсказки для анкеты: бренды каталога (по алфавиту)."""
    return sorted({n for n, _, _ in _catalog_brands().values()}, key=str.lower)


def xlsx(days=30):
    import io
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = report(days)
    wb = Workbook()
    ws = wb.active
    ws.title = "Бренды"
    ws.append([f"AMURA — пожелания клиентов за {r['days']} дн.: запрошенные бренды"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Бренд", "Просили, раз", "Есть в каталоге", "Товаров", "Остаток, шт."])
    for c in ws[2]:
        c.font = Font(bold=True)
    for t in r["brands"]:
        ws.append([t["brand"], t["count"], "да" if t["inCatalog"] else "нет", t["products"], t["qty"]])
    ws2 = wb.create_sheet("Пожелания")
    ws2.append(["Дата", "Откуда", "Что хотят", "Бренды", "Что улучшить", "Оценка", "Имя", "Телефон", "Отметка"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    st = {"new": "новое", "done": "учли", "answered": "ответили"}
    for i in r["items"]:
        ws2.append([datetime.fromtimestamp(i["at"], oh.ALMATY).strftime("%d.%m.%Y %H:%M"), i["siteName"], i["want"], ", ".join(i["brands"]),
                    i["improve"], i["rating"] or "", i["name"], i["phone"], st.get(i["status"], "")])
    for w, cols in ((ws, (28, 13, 15, 10, 13)), (ws2, (17, 18, 50, 28, 50, 8, 18, 15, 11))):
        for n, width in enumerate(cols):
            w.column_dimensions[chr(65 + n)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
